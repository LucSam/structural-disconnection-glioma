#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document as create_document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TABLE_DIR = ROOT / "outputs" / "tables"
METADATA_PATH = ROOT / "metadata" / "table_export_metadata.json"


@dataclass(frozen=True)
class TableSpec:
    number: int
    stem: str
    orientation: str
    font_size: float
    widths_inches: tuple[float, ...]


TABLES = (
    TableSpec(1, "table_01_cohort_characteristics", "portrait", 9.0, (2.15, 4.75)),
    TableSpec(
        2,
        "table_02_tfnbs_subtest_summary",
        "landscape",
        9.0,
        (0.68, 1.38, 1.05, 1.13, 1.30, 1.42, 3.70),
    ),
    TableSpec(
        3,
        "table_03_prediction_models",
        "landscape",
        9.0,
        (2.15, 2.10, 1.55, 1.55, 1.55, 1.60),
    ),
)


def load_table_metadata() -> dict[str, dict[str, str]]:
    metadata = json.loads(METADATA_PATH.read_text(encoding="utf-8"))
    for spec in TABLES:
        entry = metadata.get(str(spec.number))
        if not isinstance(entry, dict) or not entry.get("title") or not entry.get("note"):
            raise ValueError(f"Missing Table {spec.number} title or note in {METADATA_PATH}")
    # Benchmark estimates belong in the note so all table rows share six columns.
    # Read their saved results on every export, rather than hard-coding numbers.
    profile_path = ROOT / "results/analysis/prediction/demtect_profile_model_summary.csv"
    with profile_path.open(newline="", encoding="utf-8") as handle:
        profiles = {row["feature_set"]: row for row in csv.DictReader(handle)}
    regional = profiles["clinical_lobe_roi_pca"]
    baseline = profiles["clinical_lobe"]

    def interval(row: dict[str, str], field: str) -> str:
        return (f"{float(row[field]):.3f} "
                f"({float(row[field + '_ci_low']):.3f} to {float(row[field + '_ci_high']):.3f})")

    metadata["3"]["note"] = metadata["3"]["note"].format(
        training_mean=interval(regional, "fold_mean_profile_benchmark_median_profile_r"),
        baseline_vs_mean=interval(baseline, "delta_median_profile_r_vs_fold_mean_profile_benchmark"),
        regional_vs_mean=interval(regional, "delta_median_profile_r_vs_fold_mean_profile_benchmark"),
    )
    return metadata


def read_csv_rows(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if len(rows) < 2 or not rows[0]:
        raise ValueError(f"Expected a header and data rows in {path}")
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise ValueError(f"Inconsistent column count in {path}")
    return rows


def set_cell_width(cell, width_inches: float) -> None:
    width_twips = int(width_inches * 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(width_twips))
    tc_width.set(qn("w:type"), "dxa")


def set_cell_margins(cell, margin_twips: int = 65) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def configure_page(document: DocxDocument, orientation: str) -> None:
    section = document.sections[0]
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297)
        section.page_height = Mm(210)
        margin = Mm(11)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        margin = Mm(16)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin


def set_document_defaults(document: DocxDocument, font_size: float) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(font_size)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def add_title(document: DocxDocument, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)


def add_table(document: DocxDocument, rows: list[list[str]], spec: TableSpec) -> None:
    if len(rows[0]) != len(spec.widths_inches):
        raise ValueError(
            f"{spec.stem} has {len(rows[0])} columns; expected {len(spec.widths_inches)}"
        )
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, spec.widths_inches, strict=True):
        column.width = Inches(width)
    mark_header_row(table.rows[0])

    for row_index, (word_row, source_row) in enumerate(zip(table.rows, rows, strict=True)):
        for column_index, (cell, value) in enumerate(zip(word_row.cells, source_row, strict=True)):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_width(cell, spec.widths_inches[column_index])
            set_cell_margins(cell)
            if row_index == 0:
                shade_cell(cell, "E7E6E6")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(spec.font_size)
            run.bold = row_index == 0
            if row_index == 0:
                run.font.color.rgb = RGBColor(0, 0, 0)
        # Keep an estimate and its confidence interval together across pages.
        tr_pr = word_row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    if spec.number == 3:
        # Repeated outcome labels need appear only once, within the same table.
        start = 1
        while start < len(rows):
            end = start + 1
            while end < len(rows) and rows[end][0] == rows[start][0]:
                end += 1
            if end - start > 1:
                for index in range(start + 1, end):
                    table.cell(index, 0).text = ""
                merged = table.cell(start, 0).merge(table.cell(end - 1, 0))
                # Merging retains empty paragraphs; remove those only.
                for paragraph in list(merged.paragraphs)[1:]:
                    if not paragraph.text:
                        merged._tc.remove(paragraph._p)
            start = end


def add_note(document: DocxDocument, note: str, font_size: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    label = paragraph.add_run("Note. ")
    label.bold = True
    label.italic = True
    label.font.name = "Arial"
    label.font.size = Pt(max(font_size, 8.0))
    run = paragraph.add_run(note)
    run.font.name = "Arial"
    run.font.size = Pt(max(font_size, 8.0))


def export_table(spec: TableSpec, metadata: dict[str, dict[str, str]]) -> Path:
    source = TABLE_DIR / f"{spec.stem}.csv"
    destination = TABLE_DIR / f"{spec.stem}.docx"
    rows = read_csv_rows(source)
    if spec.number == 3:
        rows[0] = ["Outcome", "Measure", "Clinical only", "Baseline", "Regional", "Δ vs baseline"]
    title = metadata[str(spec.number)]["title"].strip()
    note = metadata[str(spec.number)]["note"].strip()

    document = create_document()
    configure_page(document, spec.orientation)
    set_document_defaults(document, spec.font_size)
    document.core_properties.title = title
    document.core_properties.subject = "Editable manuscript table generated from the canonical CSV"
    document.core_properties.keywords = "ICONS-GP; manuscript table; reproducible export"
    add_title(document, title)
    add_table(document, rows, spec)
    add_note(document, note, spec.font_size)
    document.save(str(destination))
    return destination


def main() -> None:
    metadata = load_table_metadata()
    exported = [export_table(spec, metadata) for spec in TABLES]
    for path in exported:
        print(f"Wrote {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
