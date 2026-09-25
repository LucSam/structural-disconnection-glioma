#!/usr/bin/env python3
"""Build the manuscript and editable tables from the checked result tables."""
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from jinja2 import Environment, StrictUndefined

from _common import ROOT, sha256


def probability(value):
    """Keep the permutation resolution without unnecessary trailing zeroes."""
    return f'{float(value):.5f}'.rstrip('0').rstrip('.')


def context_from_results():
    models = pd.read_csv(ROOT / 'results/staged_models.csv')
    groups = json.loads((ROOT / 'metadata/clinical_definitions.json').read_text())
    thirds = pd.read_csv(ROOT / 'results/burden_tertiles.csv')
    qol = pd.read_csv(ROOT / 'results/qol_correlations.csv')
    high_qol = pd.read_csv(ROOT / 'results/high_burden_preserved_qol.csv')
    context = {}
    for battery, short, outcome in [('AAT', 'aat', 'AAT_mean_T_score'),
                                    ('DemTect', 'dem', 'DemTect_global')]:
        g = groups[battery]
        context.update({f'{short}_high_preserved': g['n_high_burden_preserved'],
                        f'{short}_high_n': g['n_high_burden'],
                        f'{short}_preserved': g['n_preserved'],
                        f'{short}_high_percent': f"{100*g['n_high_burden_preserved']/g['n_high_burden']:.1f}"})
        t = thirds[thirds.battery.eq(battery)]
        values = [f'{r.preserved}/{r.n}' for r in t.itertuples()]
        context[f'{short}_tertile_preserved'] = ', '.join(values[:-1]) + ' and ' + values[-1]
        q = qol[qol.neuropsych_measure.eq(outcome) &
                qol.qol_measure.eq('EORTC_QLQ_BN20_communication_deficit')].iloc[0]
        context[f'{short}_comm_rho'] = f'{q.spearman_rho:.3f}'
        context[f'{short}_comm_q'] = f'{q.q_fdr_all_neuropsych_qol_tests:.4f}'
        counts = high_qol[high_qol.battery.eq(battery)].n_qol_available.unique()
        assert len(counts) == 1, 'Revise the manuscript wording if subgroup QoL counts differ by scale'
        context[f'{short}_high_qol_n'] = int(counts[0])
        for question, label in [('Extent beyond clinical factors', 'extent'),
                                ('Pattern beyond extent', 'pattern'),
                                ('Pattern sensitivity', 'anatomy')]:
            row = models[models.battery.eq(battery) & models.question.eq(question)].iloc[0]
            context[f'{short}_{label}_trace'] = f'{row.pillai:.3f}'
            context[f'{short}_{label}_p'] = probability(row.p_permutation)
            context[f'{short}_{label}_q'] = probability(row.q_fdr_six_models)
    high_groups = pd.read_csv(ROOT / 'results/high_burden_group_summary.csv')
    anatomy = pd.read_csv(ROOT / 'results/high_burden_anatomical_means.csv')
    for battery, short in [('AAT','aat'),('DemTect','dem')]:
        for group, tag in [('No deficit','unremarkable'),('Deficit','low')]:
            row = high_groups[high_groups.battery.eq(battery) & high_groups.group.eq(group)].iloc[0]
            context[f'{short}_high_mean_{tag}'] = f'{100*row.mean_burden:.1f}'
            for name, key in [('Temporal','temporal'),('Parietal','parietal'),('Frontal/insula','frontal')]:
                scope = 'left-labelled parcels' if battery == 'AAT' else 'all parcels'
                row = anatomy[anatomy.battery.eq(battery) & anatomy.group.eq(group) & anatomy.scope.eq(scope) & anatomy.anatomical_group.eq(name)].iloc[0]
                context[f'{short}_{key}_{tag}'] = f'{100*row.mean_chaco:.1f}'
    counts = pd.read_csv(ROOT / 'results/connection_model_comparison.csv')
    positive = counts[counts.hemisphere_across_subtests.gt(0)]
    labels = positive.subtest.tolist()
    joined = ', '.join(labels[:-1]) + ' and ' + labels[-1] if len(labels)>1 else (labels[0] if labels else '')
    context['hemisphere_results_sentence'] = (
        f'The clinical-plus-hemisphere model retained connections for {joined} after edge-FWE and across-subtest correction.'
        if len(labels) else 'The clinical-plus-hemisphere model did not retain connections after edge-FWE and across-subtest correction.')
    network = pd.read_csv(ROOT / 'results/network_qol_correlations.csv')
    for r in network[network.qol_measure.eq('EORTC_QLQ_BN20_communication_deficit')].itertuples():
        context[f'{r.feature}_comm_rho'] = f'{r.rho:.3f}'
    edge_sets = pd.read_csv(ROOT / 'results/qol_connection_sets.csv')
    differences = pd.read_csv(ROOT / 'results/high_burden_edge_group_differences.csv')
    for battery, short in [('AAT','aat'),('DemTect','dem')]:
        context[f'{short}_qol_edges'] = f'{int(edge_sets.battery.eq(battery).sum()):,}'
        d = differences[differences.battery.eq(battery)].set_index(['group_i','group_j'])
        context[f'{short}_temp_par_difference'] = f'{d.loc[("Temporal","Parietal"),"difference_percentage_points"]:.1f}'
        context[f'{short}_frontal_difference'] = f'{-d.loc[("Frontal/insula","Frontal/insula"),"difference_percentage_points"]:.1f}'
    support = pd.read_csv(ROOT / 'results/demtect_regional_support.csv')
    for deficit, tag in [(True,'deficit'),(False,'no_deficit')]:
        r = support[support.lesion_side.eq('L') & support.deficit.eq(deficit) &
                    support.roi_name.eq('S_subparietal (R)')].iloc[0]
        context[f'dem_right_subparietal_{tag}'] = f'{100*r.mean_chaco:.2f}'
    overlap = json.loads((ROOT / 'metadata/battery_overlap.json').read_text())
    context.update({f'overlap_{key}': value for key, value in overlap.items()
                    if isinstance(value, (int, float))})
    context['overlap_burden_rho'] = f"{overlap['burden_spearman_rho_paired']:.3f}"
    return context, models, groups


def manuscript_tables(models, groups):
    cohort = pd.read_csv(ROOT / 'data/reference/cohort_characteristics.csv')
    replace = {
        'Age at inclusion': ('Age, mean (SD)', '51.2 (15.8) years'),
        'Sex': ('Sex', '75 women; 88 men'),
        'WHO grade': ('WHO grade at surgery', 'I: 1; II: 33; III: 39; IV: 90'),
        'Histological entity': ('Recorded histological entity', 'Glioblastoma: 81; astrocytoma: 45; oligodendroglioma: 25; other/rare glioma entities: 12'),
        'Hemisphere': ('Lesion hemisphere', 'Left: 151; right: 12'),
        'Handedness': ('Handedness', 'Right: 140; left: 13; ambidextrous: 4; missing: 6'),
        'Tumour volume': ('Tumour volume', 'Median 22.1 ml; IQR 11.6–53.9; range 0.3–183.9'),
        'AAT raw total / mean T-score availability': ('AAT availability', 'Recorded mean T-score: 135; all four subtests: 134'),
        'Mean AAT T-score operational group': ('Mean AAT T-score <63.5', '32/135 (23.7%)'),
        'AAT subtest T-scores below boundary': ('Low AAT subtest scores', 'Token Test: 20/135; Repetition: 32/135; Naming: 40/136; Comprehension: 43/136'),
        'DemTect availability/impairment': ('DemTect global score <13', '60/132 (45.5%)'),
        'Exploratory QoL endpoints': ('Patient-reported outcome availability', 'Global health/QoL: 64; cognitive functioning: 62; communication: 70'),
    }
    rows = [replace.get(r.Variable, (r.Variable, r.Value)) for r in cohort.itertuples()]
    for battery in ['AAT', 'DemTect']:
        g = groups[battery]
        rows.extend([
            (f'{battery}: no deficit in the model sample', f"{g['n_preserved']}/{g['n']}"),
            (f'{battery}: no deficit in the highest burden third',
             f"{g['n_high_burden_preserved']}/{g['n_high_burden']} ({100*g['n_high_burden_preserved']/g['n_high_burden']:.1f}%)"),
        ])
    g = groups['AAT']
    rows.insert(11, ('AAT complete cases: score categories',
                     f"No AAT subtest deficit: {g['n_preserved']}; Subtest deficit(s) with mean ≥63.5: {g['n']-g['n_preserved']-g['n_global_below']}; mean <63.5: {g['n_global_below']} (n = {g['n']})"))
    table1 = pd.DataFrame(rows, columns=['Characteristic', 'Value'])
    modelrows = []
    for r in models.itertuples():
        comparison = {'Extent beyond clinical factors': 'Disconnection burden beyond clinical factors',
                      'Pattern beyond extent': 'Disconnection pattern beyond burden (8 PCs)',
                      'Pattern sensitivity': 'Disconnection pattern beyond burden (anatomical-group sensitivity)'}[r.question]
        modelrows.append([r.battery, comparison, str(r.n), f'{r.pillai:.3f}',
                          probability(r.p_permutation), probability(r.q_fdr_six_models), 'FDR: 6 joint tests'])
    table2 = pd.DataFrame([r[:6] for r in modelrows], columns=['Battery', 'Comparison', 'n', "Pillai’s trace", 'p', 'FDR q'])
    counts = pd.read_csv(ROOT / 'results/connection_model_comparison.csv')
    table3 = counts[['battery','subtest','n','hemisphere_edge_FWE','hemisphere_across_subtests','lobe_edge_FWE','lobe_across_subtests']].copy()
    for col in table3.columns[3:]:
        table3[col] = table3[col].map(lambda value: f'{int(value):,}' if value else '—')
    table3['n'] = table3.n.astype(str)
    table3.columns = ['Battery','Subtest','n','Clinical + hemisphere: Edge FWE','Clinical + hemisphere: Across subtests',
                      '+ Lobe: Edge FWE','+ Lobe: Across subtests']
    return [
        {'stem': 'table_01_cohort', 'title': 'Table 1. Cohort, assessment availability and clinical description.',
         'frame': table1, 'widths': [6.4, 10.5],
         'legend': 'Values are counts unless stated otherwise. WHO grade and histological entity reflect the diagnosis recorded at surgery. AAT availability varies by subtest; complete four-subtest cases define the model sample and burden thirds. No AAT deficit requires Token Test ≥63, Repetition ≥63, Naming ≥64 and Comprehension ≥64. No DemTect deficit requires a global score ≥13. The highest third is defined within each model sample using disconnection burden: mean ChaCo over 82 left cerebral and ten left cerebellar parcels for AAT, and over all 191 parcels for DemTect. QoL availability is shown before matching to individual test measures. IQR, interquartile range; SD, standard deviation.'},
        {'stem': 'table_03_staged_models', 'title': 'Table 3. Associations with disconnection burden and pattern.',
         'frame': table2, 'widths': [1.8, 6.5, 1.1, 2.1, 2.65, 2.65],
         'legend': 'All rows test the four AAT or five DemTect subtests together. Disconnection burden is added to age, WHO grade, log-transformed tumour volume and hemisphere. The disconnection pattern is then added as eight regional principal components (PCs). Anatomical sensitivity models replace PCs with eight AAT or seven nonredundant DemTect group means. Disconnection burden is mean ChaCo over 82 left cerebral and ten left cerebellar parcels for AAT, and all 191 parcels for DemTect. FDR q values cover these six comparisons. All use 19,999 Freedman–Lane permutations. Pillai’s trace is a multivariate statistic, not a percentage of explained performance variance. Profile checks are reported separately in Supplementary Tables S1–S2.'},

        {'stem': 'table_04_connection_models', 'title': 'Table 4. Connection associations across all subtests and both anatomical models.',
         'frame': table3, 'widths': [1.95,3.6,1.0,2.6,2.7,2.4,2.6],
         'legend': 'Entries are numbers of connections; a dash indicates zero. The starting model includes age, WHO grade, tumour volume and hemisphere. The sensitivity model additionally includes frontal, temporal, parietal, occipital, insular and thalamic lesion percentages. Edge FWE denotes TFNBS family-wise error correction across connections within one subtest. Across subtests additionally applies Bonferroni correction across four AAT or five DemTect subtests. Both models use the same patients for each endpoint and 5,000 shuffles. The model comparison describes sensitivity to coarse lesion location; it is not a test of the difference between subtests or independent replication.'},
    ]


def statistical_overview():
    frame = pd.DataFrame([
        ['Do disconnection burden or pattern add information?', 'All subtests together, separately for AAT and DemTect; anatomical-group sensitivity', 'FDR across 6 joint tests', 'Central analysis and robustness'],
        ['Do associations differ between subtests?', 'Joint profile contrasts', 'Holm across 5 profile tests', 'Interpretation check; Table S1'],
        ['Which connections relate to lower scores?', 'Clinical + hemisphere TFNBS; additional lobe adjustment', 'FWE across edges, then Bonferroni across 4 AAT or 5 DemTect subtests; separately per model', 'Anatomical localisation and sensitivity'],
        ['Does measured performance relate to reported function?', 'Unadjusted Spearman correlations', 'FDR across all 222 estimable correlations', 'Patient perspective'],
        ['How does disconnection relate to reported function?', 'Disconnection burden and selected-connection summaries versus QoL', 'Descriptive coefficients; regional tests: FDR across 6 and 51 pairs (Table S3)', 'Structural–QoL description'],
        ['Where do disconnection and performance vary?', 'Regional and connection coverage, group means and differences; regional partial correlations', 'No tests or significance thresholds', 'Descriptive anatomical context'],
    ], columns=['Question', 'Analysis', 'Correction', 'Role'])
    frame = frame.iloc[[5,0,1,2,3,4]].reset_index(drop=True)
    return {'stem':'table_02_analysis_overview','title':'Table 2. Questions, analyses and correction families.',
            'frame':frame,'widths':[3.9,4.3,5.4,3.3],
            'legend':'The central model comparisons address disconnection burden and pattern. The localising analyses and patient-reported outcomes answer complementary questions; their p values are not combined into an overall test. Pairwise profile checks and their separate 16-test Holm correction are documented in Table S2; their interpretation required a significant parent profile test. FDR, false discovery rate; FWE, family-wise error. The hierarchy describes an exploratory revision, not prospectively registered analyses.'}


def supplementary_tables():
    profile = pd.read_csv(ROOT / 'results/profile/profile_omnibus_tests.csv')
    rows=[]
    for r in profile.itertuples():
        label = 'Original AAT T-scores' if r.scale == 'original_T_scores' else ('8 PCs' if r.representation == 'pc8' else 'Anatomical means')
        rows.append([r.battery,label,str(r.n),f'{r.pillai:.3f}',probability(r.p_permutation),probability(r.p_holm_5)])
    profile_frame = pd.DataFrame(rows,columns=['Battery','Profile model','n',"Pillai’s trace",'p','Holm p'])
    pairs = pd.read_csv(ROOT / 'results/profile/pairwise_profile_tests.csv')
    pair_frame = pd.DataFrame([[r.battery,r.contrast,str(r.n),probability(r.p_permutation),probability(r.p_holm_16)]
                              for r in pairs.itertuples()],columns=['Battery','Subtest contrast','n','p','Holm p'])
    prior = pd.read_csv(ROOT / 'results/prior_regional_qol_families.csv')
    prior_frame = pd.DataFrame([
        ['Whole-brain burden' if r.feature == 'whole_mean' else 'Top-decile ChaCo',
         r.qol_label,str(r.n),f'{r.rho:.3f}',probability(r.p),probability(r.q_six),probability(r.q_joint_51)]
        for r in prior.itertuples()],
        columns=['Regional ChaCo','QoL domain','n','Spearman ρ','p','FDR q: 6','FDR q: 51'])
    return [
        {'stem':'table_S1_profile_tests','title':'Table S1. Joint profile checks.', 'frame':profile_frame,
         'widths':[1.8,6.5,1.1,2.1,2.65,2.65],
         'legend':'Five tests share one Holm correction. The two PC and two anatomical models use inverse-normal rank-transformed outcomes. The original-T-score variant is an AAT sensitivity check. All use 19,999 Freedman–Lane permutations and retain the central model covariates and disconnection burden. None survived correction.'},
        {'stem':'table_S2_pairwise_profiles','title':'Table S2. Pairwise subtest profile contrasts.', 'frame':pair_frame,
         'widths':[1.8,8.3,1.1,2.8,2.8],
         'legend':'All 16 contrasts share one Holm correction and use the PC pattern model with 19,999 permutations. Interpretation required the corresponding parent PC profile test to survive the five-test correction in Table S1. Neither parent did, and no pairwise contrast survived the 16-test correction. These results are reported for completeness and were not used to select maps.'},
        {'stem':'table_S3_regional_qol_families','title':'Table S3. Regional disconnection–QoL correlations and correction families.',
         'frame':prior_frame,'widths':[2.6,3.9,.8,2.0,2.2,2.2,2.2],
         'legend':'Unadjusted Spearman correlations; higher scores indicate better reported function. Whole-brain burden is mean ChaCo across all 191 parcels; top-decile ChaCo is the mean of the largest 20 values per patient. The six p values have FDR correction across these six pairs and, separately, across all 51 original structural tests (these six plus 45 graph-descriptor correlations). All six survive the narrower correction; none survives the broader correction. Fixed formulas do not establish prospective specification of the six-test family. The broader sensitivity correction was post hoc. Figure 5 additionally displays descriptive correlations for left-hemisphere burden and selected edges, which were not tested within these historical families.'},
    ]


FIGURES = [
    ('figure_01_cohort_coverage', 'Figure 1. Tumour distribution, estimated disconnection and clinical performance.',
     'A: Lesion overlap across all 163 patients. B: Mean regional ChaCo across the cohort. C: Mean ChaCoConn by anatomical endpoints. All 18,145 unordered atlas pairs enter the group-pair means, including zeros and both hemispheres; diagonal cells contain connections between distinct parcels within a group. D: Recorded mean AAT T-score, with the operational boundary at 63.5. E: Principal AAT subtest T-scores; lower boundaries are 63 for Token Test and Repetition and 64 for Naming and Comprehension. F: DemTect global score, with its screening boundary at 13. G: DemTect subtest scores as percentages of their maxima; the global-score cutoff does not apply to these subtests. Boxplots show medians, interquartile ranges and whiskers to the most extreme observations within 1.5 interquartile ranges; points are individual observations with vertical jitter. Figure S1 illustrates how the regional and connection estimates are obtained.'),
    ('figure_02_disconnection_performance', 'Figure 2. Regional and connection differences at high burden, with individual performance.',
     'A–D: Regional and connection differences (deficit minus no deficit) within each battery’s highest burden third. Red indicates greater disconnection with deficits; blue, without. All 191 parcels and 18,145 pairs are included; connection cells average pairs by anatomical group, including both hemispheres and zeros. Separate scales use percentage points. Groups are unmatched and unadjusted; means are in Figures S2–S3. E–F: Individual performance with third boundaries and deficit counts. Outlines mark high burden without deficits. AAT deficit means any of four subtests below its cutoff; the plot shows the lowest subtest T-score minus its cutoff (≥0: all four met). DemTect deficit means global score <13. Burden averages 82 left cerebral and ten left cerebellar parcels for AAT, all 191 for DemTect. Colours distinguish no deficit, AAT subtest deficits with mean ≥63.5, and low global scores. Small vertical jitter separates AAT values tied at +8; original values define groups.'),
    ('figure_03_regional_associations', 'Figure 3. Regional disconnection associations with overall test performance.',
     'Unthresholded partial deficit correlations for mean AAT T-score (A) and DemTect global score (B), adjusted for age, WHO grade, tumour volume and hemisphere. Positive values indicate greater disconnection with lower performance. Each map displays all 191 atlas parcels on a common scale, in left lateral, coronal, right lateral and axial projection. Signed projections retain the value of greatest absolute magnitude along each viewing ray. Colours describe coefficients, not significance; the nine principal-subtest maps are in Figure S4.'),
    ('figure_04_connection_anatomy', 'Figure 4. Connection associations after clinical and hemisphere adjustment.',
     'All connections surviving edge-FWE and across-subtest correction are shown, with no numerical cap or strongest-edge selection. Panels include every subtest with retained connections; Table 4 reports all nine. Parcels are ordered in mirrored anatomical groups on left and right halves, with vermis below. Edge colour, width and transparency increase with the edge-wise t statistic; node colours identify anatomical groups. Each connection is drawn as a separate curved line. Curves join atlas endpoints and do not depict fibre trajectories. Larger nodes have at least one retained connection. The panels localise associations and do not directly test differences between subtests. Lobe-adjusted maps are shown separately in Figure S5.'),
    ('figure_05_patient_reported_function', 'Figure 5. Performance, disconnection and patient-reported function.',
     'A: Spearman correlations for 11 displayed neuropsychological measures and three QoL domains. Higher QoL values indicate better status, including the reversed communication scale. Asterisks denote FDR q <0.05 across all 222 estimable test–QoL pairs; the undisplayed AAT raw total remains in that family. B: Descriptive structural–QoL coefficients, without significance markers. Left-hemisphere burden uses mean ChaCo over 82 left cerebral and ten left cerebellar parcels; whole-brain burden uses all 191 parcels. AAT and DemTect edge summaries average the unique main-model connections surviving both TFNBS correction steps, pooled across subtests within each battery (2,930 and 404 edges). Selection used performance in this cohort. C–F: Communication versus performance or burden. Grey LOWESS guides in C–D are displayed within 0–100. E–F show individual observations without a fitted curve; the reported Spearman coefficients use all paired observations. Colours follow the corresponding battery; grey points lack complete test classification. Outlines mark high burden without deficits (eight AAT, six DemTect cases). Structural–QoL pairs include patients with missing test data. Performance plots have horizontal jitter up to 0.1; all plots have vertical jitter up to 0.5. Curves and statistics use original values. No separate subgroup tests were fitted.'),
]

SUPPLEMENT_FIGURES = [
    ('figure_S4_subtest_regional_associations', 'Figure S4. Regional associations for all principal subtests.',
     'Partial deficit correlations adjusted for age, WHO grade, tumour volume and hemisphere, using the same definition as Figure 3. Each panel shows all 191 parcels without thresholding, on one common scale, in left lateral, axial and right lateral projection. Positive values indicate greater disconnection with lower performance. The maps are descriptive and do not establish different effects between subtests.'),
    ('figure_S5_lobe_connection_anatomy', 'Figure S5. Connection maps with additional lobe adjustment.',
     'All connections surviving edge-FWE and across-subtest correction in the additionally lobe-adjusted model are shown, without a numerical cap. Only Token Test (21) and Naming (152) meet both steps; Table 4 reports all nine subtests and both correction stages. Colours and individual connection arcs follow Figure 4, with the same t-statistic scale. These panels use the same display threshold as Figure 4 and add lobar covariates. Each model is corrected separately, so these connections need not be a subset of those in Figure 4. Full edge tables identify the connections meeting each correction step.'),
    ('figure_S3_group_connection_disconnection', 'Figure S3. Anatomical connection summaries in high-burden groups.',
     'Each row shows patients without a deficit, patients with a deficit, and the deficit-minus-no-deficit difference within the highest burden third. AAT groups require complete four-subtest data; DemTect groups use the global-score boundary of 13. All mean panels share one scale, and both difference panels share a symmetric scale in percentage points. Cells average all atlas-pair ChaCoConn estimates within anatomical group pairs, including zero values and both hemispheres, with equal weight per pair. Within-group diagonal cells exclude self-connections. All 18,145 unordered atlas pairs are represented, without selection by TFNBS or performance association. These are unadjusted, unmatched descriptions; positive or negative differences do not establish vulnerability or protection.'),
    ('figure_S1_nemo_estimation', 'Figure S1. How NeMo estimates regional and connection disconnection.',
     'A synthetic six-region example on a Nilearn glass-brain outline. A: Each straight line summarises one reference connection with four units of streamline weight; the shaded area represents a tumour mask. B: Affected weights of one, two, three or four units give pairwise losses of 25%, 50%, 75% or 100% (ChaCoConn). Colour and positive line width encode this percentage; dashed grey lines indicate 0% loss. C: Node colour shows regional ChaCo, calculated by dividing the total affected weight at a node by its total reference weight. For A, this is (3 + 1)/12 = 33.3%. Lines retain the pairwise values from B; node labels are rounded to one decimal place. D: The matrix contains the same pairwise losses. Grey cells have no reference connection and are undefined. The illustrative affected weights are assigned, not calculated from intersections with the two-dimensional mask; straight lines connect endpoints and do not depict fibre trajectories. The study estimates these fractions from three-dimensional reference streamlines using 191 parcels and averages them over healthy reference connectomes.'),
    ('figure_S2_regional_group_means', 'Figure S2. Regional disconnection in high-burden groups.',
     'All 191 regional ChaCo means are shown for patients without and with deficits, followed by deficit-minus-no-deficit differences. Rows use the same battery-specific highest-third groups as Figure 2. All mean maps share one scale and both difference maps share a symmetric percentage-point scale. AAT disconnection burden averages 82 left cerebral and ten left cerebellar parcels; DemTect disconnection burden averages all 191. The maps show all parcels for both batteries. Atlas-filled projections retain the greatest absolute value along each viewing ray, so parcels may overlap. These are unadjusted, unmatched descriptions.'),
]

SUPPLEMENT_FIGURES.sort(key=lambda item: int(re.search(r'figure_S(\d+)',item[0]).group(1)))



def add_inline(paragraph, text):
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        run = paragraph.add_run(part)
        run.bold = bool(i % 2)


def style_document(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    section.left_margin, section.right_margin = Cm(2.05), Cm(2.05)
    section.header_distance, section.footer_distance = Cm(0.8), Cm(0.8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'List Bullet']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().append(OxmlElement('w:lang'))
        style._element.rPr[-1].set(qn('w:val'), 'en-GB')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
    doc.styles['Caption'].font.size = Pt(10)
    doc.styles['Caption'].font.italic = False
    doc.styles['Caption'].font.bold = False
    title_properties = doc.styles['Title']._element.pPr
    if title_properties is not None:
        for border in list(title_properties.findall(qn('w:pBdr'))):
            title_properties.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run()
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)
    doc.core_properties.author = 'Lucius S. Fekonja'
    doc.core_properties.subject = 'Structural disconnection, preoperative performance and patient-reported function in glioma'
    doc.core_properties.keywords = 'glioma; NeMo; disconnection; AAT; DemTect; quality of life'


def add_table(doc, spec):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.add_run(spec['title']).bold = True
    df = spec['frame']
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'
    for col, width in zip(table.columns, spec['widths'], strict=True):
        col.width = Cm(width)
    headers = list(df.columns)
    for cell, header, width in zip(table.rows[0].cells, headers, spec['widths'], strict=True):
        cell.width = Cm(width)
        cell.text = header
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E8EDF0')
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in df.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for j, (cell, value, width) in enumerate(zip(cells, values, spec['widths'], strict=True)):
            cell.width = Cm(width)
            cell.text = str(value)
    for row in table.rows:
        cant_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant_split)
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.line_spacing = 1.0
                if len(df.columns) > 2 and j >= 2 and spec['stem'] != 'table_02_analysis_overview':
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9.5)
    p = doc.add_paragraph(spec['legend'])
    p.paragraph_format.space_before = Pt(8)
    for run in p.runs:
        run.font.size = Pt(9.5)


def selected_references(body):
    library = json.loads((ROOT / 'manuscript/reference_library.json').read_text())
    selected = []
    for ref in library:
        author, year = ref['key'].rsplit(' ', 1)
        if re.search(re.escape(author) + r'[^();\n]{0,65}\b' + year + r'\b', body):
            text = re.sub(r'(?<=\w)-\s+(?=\w)', '-', ref['text'])
            selected.append({'key': ref['key'], 'text': text})
    return sorted(selected, key=lambda r: re.sub(r'[^a-z0-9]', '', r['key'].lower()))


def table_markdown(spec):
    df=spec['frame']
    lines=['### '+spec['title'], '', '| '+' | '.join(df.columns)+' |',
           '| '+' | '.join(['---']*len(df.columns))+' |']
    lines.extend('| '+' | '.join(map(str,row))+' |' for row in df.itertuples(index=False,name=None))
    return '\n'.join(lines)+'\n\n'+spec['legend']


def render_body(doc, body, overview=None):
    section='title'
    declarations={'Acknowledgements','Contributions','Corresponding author',
                  'Competing interests','Data availability'}
    for block in body.split('\n\n'):
        block=block.strip()
        if not block:
            continue
        if block in ['## Abstract','## 1 Introduction']:
            doc.add_page_break()
        if block == '[[STATISTICAL_OVERVIEW]]':
            add_table(doc,overview)
        elif block.startswith('# '):
            doc.add_paragraph(block[2:], 'Title')
        elif block.startswith('### '):
            doc.add_heading(block[4:],level=2)
        elif block.startswith('## '):
            section=block[3:]
            heading=doc.add_heading(section,level=1)
            if section in declarations:
                heading.paragraph_format.space_before=Pt(6)
                heading.paragraph_format.space_after=Pt(3)
                for run in heading.runs:
                    run.font.size=Pt(12)
        elif block.startswith('- '):
            for line in block.splitlines():
                add_inline(doc.add_paragraph(style='List Bullet'),line.removeprefix('- '))
        else:
            p=doc.add_paragraph();add_inline(p,block)
            if section == 'title':
                for run in p.runs:
                    run.font.size=Pt(10 if block.startswith(('¹','²','³','⁴')) else 11)
            elif section in declarations:
                p.paragraph_format.line_spacing=1.
                p.paragraph_format.space_after=Pt(4)
                for run in p.runs:
                    run.font.size=Pt(10)


def add_references(doc, refs, title):
    doc.add_heading(title,level=1)
    for ref in refs:
        p=doc.add_paragraph(ref['text'])
        p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5)
        p.paragraph_format.line_spacing=1.;p.paragraph_format.keep_together=True
        for run in p.runs:run.font.size=Pt(10)


def add_figures(doc, figures):
    md=''
    for stem,title,caption in figures:
        p=doc.add_paragraph();p.paragraph_format.page_break_before=True
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/'outputs/figures'/f'{stem}.png'),width=Cm(16.8))
        p=doc.add_paragraph(style='Caption');p.add_run(title+' ').bold=True;p.add_run(caption)
        md+=f'\n\n### {title}\n\n![{title}](../outputs/figures/{stem}.png)\n\n{caption}'
    return md


def main():
    context,models,groups=context_from_results()
    env=Environment(undefined=StrictUndefined,autoescape=False)
    body=env.from_string((ROOT/'manuscript/manuscript_template.md').read_text()).render(context)
    supplement=env.from_string((ROOT/'manuscript/supplement_template.md').read_text()).render(context)
    refs=selected_references(body);supp_refs=selected_references(supplement)
    specs=manuscript_tables(models,groups);overview=statistical_overview();supp_specs=supplementary_tables()
    doc=Document();style_document(doc)
    doc.core_properties.title=body.splitlines()[0].removeprefix('# ')
    render_body(doc,body,overview)
    doc.add_page_break();add_references(doc,refs,'References')
    md=body.replace('[[STATISTICAL_OVERVIEW]]',table_markdown(overview))
    md+='\n\n## References\n\n'+'\n\n'.join(r['text'] for r in refs)
    for spec in specs:
        doc.add_page_break();add_table(doc,spec);md+='\n\n'+table_markdown(spec)
    md+=add_figures(doc,FIGURES)
    doc.add_page_break();render_body(doc,supplement)
    md+='\n\n'+supplement
    add_references(doc,supp_refs,'Supplementary references')
    md+='\n\n## Supplementary references\n\n'+'\n\n'.join(r['text'] for r in supp_refs)
    for spec in supp_specs:
        doc.add_page_break();add_table(doc,spec);md+='\n\n'+table_markdown(spec)
    md+=add_figures(doc,SUPPLEMENT_FIGURES)
    for spec in [overview,*specs,*supp_specs]:
        spec['frame'].to_csv(ROOT/'outputs/tables'/f"{spec['stem']}.csv",index=False)
        separate=Document();style_document(separate);add_table(separate,spec)
        separate.save(ROOT/'outputs/tables'/f"{spec['stem']}.docx")
    output=ROOT/'manuscript/icons_gp_manuscript_v2.docx';doc.save(output)
    (ROOT/'manuscript/icons_gp_manuscript_v2.md').write_text(md+'\n')
    (ROOT/'manuscript/references_used.json').write_text(json.dumps(refs,ensure_ascii=False,indent=2)+'\n')
    abstract=body.split('## Abstract\n',1)[1].split('**Keywords:**')[0]
    main_text=body.split('## 1 Introduction',1)[1].split('## Acknowledgements')[0]
    metadata={'template_sha256':sha256(ROOT/'manuscript/manuscript_template.md'),
              'supplement_template_sha256':sha256(ROOT/'manuscript/supplement_template.md'),
              'context':context,'abstract_words':len(abstract.split()),'main_text_words':len(main_text.split()),
              'methods_words':len(body.split('## 2 Methods')[1].split('## 3 Results')[0].split()),
              'results_words':len(body.split('## 3 Results')[1].split('## 4 Discussion')[0].split()),
              'supplement_words':len(supplement.split()),'references':len(refs),
              'tables':len(specs)+1,'supplementary_tables':len(supp_specs),'figures':len(FIGURES),
              'supplementary_figures':len(SUPPLEMENT_FIGURES),
              'result_hashes':{str(p.relative_to(ROOT)):sha256(p) for p in (ROOT/'results').rglob('*')
                               if p.is_file() and (p.name.endswith('.csv') or p.name.endswith('.csv.gz'))},
              'figure_hashes':{str((ROOT/'outputs/figures'/f'{stem}.png').relative_to(ROOT)):
                               sha256(ROOT/'outputs/figures'/f'{stem}.png') for stem,_,_ in FIGURES+SUPPLEMENT_FIGURES}}
    (ROOT/'metadata/manuscript_build.json').write_text(json.dumps(metadata,indent=2)+'\n')
    print(f'Built manuscript: {metadata["main_text_words"]} main words; {metadata["methods_words"]} Methods; '
          f'{metadata["results_words"]} Results; {len(doc.tables)} editable tables; {len(doc.inline_shapes)} figures including supplement.')


if __name__=='__main__':
    main()
