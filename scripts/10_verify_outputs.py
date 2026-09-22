#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import json
import pickle
import warnings
from pathlib import Path

import numpy as np
import pandas as pd  # type: ignore[import-untyped]
from PIL import Image  # type: ignore[import-untyped]
from scipy.io import loadmat  # type: ignore[import-untyped]

ROOT = Path(__file__).resolve().parents[1]


def require(condition: bool, message: str) -> None:
    if not condition:
        raise SystemExit(message)


def read_csv(relative_path: str) -> pd.DataFrame:
    path = ROOT / relative_path
    require(path.exists(), f"Missing computed result: {relative_path}")
    return pd.read_csv(path)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def bh_fdr(p_values: np.ndarray) -> np.ndarray:
    values = np.asarray(p_values, dtype=float)
    order = np.argsort(values)
    ranked = values[order]
    adjusted = np.minimum.accumulate((ranked * len(values) / np.arange(1, len(values) + 1))[::-1])[::-1]
    output = np.empty_like(adjusted)
    output[order] = np.minimum(adjusted, 1.0)
    return output


def assert_no_terms(frame: pd.DataFrame, terms: tuple[str, ...], label: str) -> None:
    searchable = " ".join(frame.columns.astype(str).tolist()) + "\n" + frame.to_csv(index=False)
    lowered = searchable.lower()
    found = [term for term in terms if term in lowered]
    require(not found, f"{label} contains retired terms: {found}")


def verify_nemo_encoding() -> None:
    subject_dirs = sorted((ROOT / "data/nemo/ifod2act_fs191").glob("sub-P*"))
    require(len(subject_dirs) == 163, f"Expected 163 extracted NeMo subject directories, observed {len(subject_dirs)}")
    sc_path = next(path for directory in subject_dirs for path in directory.glob("*nemoSC_sift2_mean.mat"))
    sc_raw = np.asarray(loadmat(sc_path)["SC"], dtype=float)
    require(sc_raw.shape == (191, 191), f"Unexpected nemoSC shape: {sc_raw.shape}")
    require(np.allclose(np.tril(sc_raw, -1), 0.0, atol=1e-12), "nemoSC input is not upper triangular")
    mirrored = np.triu(sc_raw, 1) + np.triu(sc_raw, 1).T
    require(np.isclose(mirrored.sum(), 2.0 * np.triu(sc_raw, 1).sum()), "nemoSC mirroring changed edge magnitude")

    cc_path = next(path for directory in subject_dirs for path in directory.glob("*chacoconn_fs191subj_mean.pkl"))
    with cc_path.open("rb") as handle, warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Please import `csr_matrix`", category=DeprecationWarning)
        cc_raw = pickle.load(handle).toarray().astype(float)
    require(cc_raw.shape == (191, 191), f"Unexpected ChaCoConn shape: {cc_raw.shape}")
    require(np.allclose(np.tril(cc_raw, -1), 0.0, atol=1e-12), "ChaCoConn input is not upper triangular")
    require(float(np.nanmin(cc_raw)) >= -1e-8 and float(np.nanmax(cc_raw)) <= 1.0 + 1e-8, "ChaCoConn is outside [0, 1]")


def verify_regional_statistics() -> None:
    global_tests = read_csv("results/analysis/regional_chaco/regional_chaco_global_score_partial_spearman.csv")
    require(global_tests.shape[0] == 382, f"Expected 382 adjusted global parcel tests, observed {len(global_tests)}")
    require(global_tests.groupby("outcome").size().to_dict() == {"AAT_total": 191, "DemTect_global": 191}, "Global parcel family is incomplete")
    require(set(global_tests["method"]) == {"partial_spearman_pearson_on_rank_residuals"}, "Regional partial Spearman method is inconsistent")
    require((global_tests["dof_partial"] > 0).all(), "Regional partial Spearman contains invalid degrees of freedom")
    expected_q = bh_fdr(global_tests["p_partial"].to_numpy())
    require(np.allclose(expected_q, global_tests["q_partial_global_score_family"], atol=1e-12), "Global parcel q values are not joint BH across 382 tests")
    observed = global_tests.groupby("outcome")["significant_global_score_family"].sum().astype(int).to_dict()
    require(observed == {"AAT_total": 7, "DemTect_global": 0}, f"Unexpected corrected global parcel counts: {observed}")
    sensitivity = read_csv("results/analysis/regional_chaco/regional_chaco_global_score_partial_sensitivity.csv")
    require(len(sensitivity) == 2 * 382, "Regional laterality sensitivity table is incomplete")
    sensitivity_counts = (
        sensitivity.groupby(["sensitivity", "outcome"])["significant_joint_382"].sum().astype(int).to_dict()
    )
    require(
        sensitivity_counts
        == {
            ("hemisphere_adjusted_all_subjects", "AAT_total"): 0,
            ("hemisphere_adjusted_all_subjects", "DemTect_global"): 0,
            ("left_hemisphere_only", "AAT_total"): 0,
            ("left_hemisphere_only", "DemTect_global"): 0,
        },
        f"Unexpected regional laterality sensitivity counts: {sensitivity_counts}",
    )

    all_tests = read_csv("results/analysis/regional_chaco/regional_chaco_partial_spearman.csv")
    summary = read_csv("results/analysis/regional_chaco/regional_chaco_partial_summary.csv")
    require(len(all_tests) == 17 * 191, "Regional outcome-by-parcel table is incomplete")
    counts = all_tests.groupby("outcome")["significant_partial_fdr"].sum().astype(int)
    reported = summary.set_index("outcome")["n_fdr_significant"].astype(int)
    require(counts.sort_index().equals(reported.sort_index()), "Regional summary counts do not match parcel-level flags")


def verify_tfnbs() -> None:
    inventory = read_csv("results/analysis/tfnbs/tfnbs_design_contrast_inventory.csv")
    counts = read_csv("results/analysis/tfnbs/tfnbs_edge_counts_from_mrtrix.csv")
    table = read_csv("results/analysis/tfnbs/table_02_tfnbs_subtest_summary.csv")
    require(len(inventory) == 18 and len(counts) == 18, "TFNBS must contain 9 subtests x 2 model families")
    require(len(table) == 9, "Table 2 must contain four AAT and five DemTect rows")
    require(len(table.columns) == 7, "Table 2 must omit the redundant narrative result column")
    for row in table.to_dict("records"):
        selected = counts[counts["battery"].eq(row["Battery"]) & counts["subtest_label"].eq(row["Subtest"])].set_index("model_family")
        for column, family, field in [
            ("Clinical: edge FWE", "clinical", "n_fwe05_edges"),
            ("Clinical: across subtests", "clinical", "n_battery_bonf05_edges"),
            ("Location-adjusted: edge FWE", "lobe", "n_fwe05_edges"),
            ("Location-adjusted: across subtests", "lobe", "n_battery_bonf05_edges"),
        ]:
            require(int(str(row[column]).replace(",", "")) == int(selected.loc[family, field]), f"Table 2 count mismatch: {row['Subtest']}/{column}")
    require(table.equals(read_csv("outputs/tables/table_02_tfnbs_subtest_summary.csv")), "Exported Table 2 differs from the analysis table")
    require(set(inventory["model_family"]) == {"clinical", "lobe"}, "Unexpected TFNBS model families")
    require({"permutation_seed_metadata_file", "permutation_seed", "permutation_seed_specification_sha256"}.issubset(inventory.columns), "TFNBS inventory omits deterministic seed provenance")

    for row in inventory.itertuples(index=False):
        design = np.loadtxt(ROOT / row.design_file, dtype=float, ndmin=2)
        contrast = np.loadtxt(ROOT / row.contrast_file, dtype=float, ndmin=2)
        require(design.shape == (int(row.n_subjects), int(row.n_design_columns)), f"TFNBS design shape mismatch: {row.model_dir}/{row.subtest_slug}")
        require(np.linalg.matrix_rank(design) == design.shape[1], f"Rank-deficient TFNBS design: {row.model_dir}/{row.subtest_slug}")
        require(np.isfinite(np.linalg.cond(design)) and np.linalg.cond(design) < 1e4, f"Ill-conditioned TFNBS design: {row.model_dir}/{row.subtest_slug}")
        expected_contrast = np.zeros((1, design.shape[1]), dtype=float)
        expected_contrast[0, 1] = -1.0
        require(np.array_equal(contrast, expected_contrast), f"Unexpected TFNBS contrast: {row.model_dir}/{row.subtest_slug}")
        if row.model_family == "lobe":
            require(
                "hemisphere_R_rank_z" in str(row.design_columns),
                f"Lobe TFNBS design omits hemisphere_R: {row.model_dir}/{row.subtest_slug}",
            )
            require(
                "lobe_hemisphere" in str(row.model_dir),
                f"Lobe TFNBS output directory is not laterality-qualified: {row.model_dir}",
            )

    output_root = ROOT / "results/analysis/tfnbs/mrtrix_outputs"
    seed_files = sorted(output_root.glob("*/*/permutation_seed.json"))
    require(len(seed_files) == 18, f"Expected 18 TFNBS seed records, observed {len(seed_files)}")
    seeds: list[int] = []
    for seed_path in seed_files:
        metadata = json.loads(seed_path.read_text(encoding="utf-8"))
        require(metadata["n_permutations"] == 5000, f"Non-release shuffle count in {seed_path.relative_to(ROOT)}")
        seeds.append(int(metadata["derived_seed"]))
        run_dir = seed_path.parent
        require((run_dir / "tfnbs_fwe_1mpvalue.csv").exists(), f"Incomplete TFNBS run: {run_dir.relative_to(ROOT)}")
        require((run_dir / "connectomestats.log").exists(), f"Missing TFNBS log: {run_dir.relative_to(ROOT)}")
    require(len(set(seeds)) == 18, "TFNBS derived seeds are not unique")


def verify_graph_statistics() -> None:
    metrics = read_csv("results/analysis/graph_topology/graph_global_metrics_from_nemo.csv")
    require(metrics.shape == (163, 16), f"Unexpected graph metric table shape: {metrics.shape}")
    require(metrics["subject_id"].is_unique, "Graph metric subject IDs are not unique")
    assert_no_terms(metrics, ("normative", "connectivity_loss", "hub_vulnerability", "sc_effective", "_delta"), "Graph metrics")
    require(
        all(
            column == "subject_id"
            or column.startswith("remaining_")
            or "chacoconn" in column
            for column in metrics.columns
        ),
        "Stable graph metric IDs do not distinguish nemoSC topology from ChaCoConn disconnection",
    )

    mode = read_csv("results/analysis/graph_topology/graph_metric_computation_mode.csv")
    description = " ".join(mode.astype(str).to_numpy().ravel()).lower()
    require("upper triangle mirrored as a + a.t" in description, "Graph provenance does not record unscaled mirroring")
    require(
        "not reapplied" in description
        and "nemosc" in description,
        "Graph provenance does not prevent double lesion attenuation or define nemoSC clearly",
    )

    binary = read_csv("results/analysis/graph_topology/graph_metric_binary_endpoint_covariate_adjusted.csv")
    binary_permutation = read_csv(
        "results/analysis/graph_topology/graph_metric_binary_endpoint_rank_freedman_lane_permutation.csv"
    )
    continuous = read_csv("results/analysis/graph_topology/graph_metric_subtest_score_partial_correlations.csv")
    nodal = read_csv("results/analysis/graph_topology/graph_nodal_outcome_correlations.csv")
    nodal_counts = read_csv("results/analysis/graph_topology/graph_nodal_signature_counts.csv")
    require(len(binary) == 2 * 15 and set(binary["fdr_family"]) == {"all_15_global_metrics_within_endpoint"}, "Binary graph FDR family is inconsistent")
    require(
        set(continuous["fdr_family_all_outcomes_check"])
        == {"all_11_outcomes_x_15_global_metrics"},
        "Continuous graph output omits the all-outcome multiplicity check",
    )
    all_outcome_significant = continuous.loc[
        continuous["significant_partial_all_165_tests"].astype(bool)
    ]
    require(
        len(all_outcome_significant) == 4
        and set(all_outcome_significant["outcome_label"]) == {"Verbal Fluency"},
        "Unexpected continuous graph findings after correction across all 165 tests",
    )
    require(
        set(binary["method"])
        == {"fully_ranked_ancova_hc3_group_plus_age_grade_log_volume_hemisphere"},
        "Binary graph primary inference is not the fully ranked HC3 ANCOVA",
    )
    require(
        len(binary_permutation) == 2 * 15
        and set(binary_permutation["n_permutations"]) == {5000}
        and set(binary_permutation["method"])
        == {"fully_ranked_freedman_lane_group_plus_age_grade_log_volume_hemisphere"},
        "Binary graph permutation check is incomplete",
    )
    expected_binary_counts = {"AAT_mean_T_below_threshold": 7, "DemTect_impaired_binary": 0}
    require(
        binary.groupby("endpoint")["significant_fdr"].sum().astype(int).to_dict()
        == expected_binary_counts,
        "Unexpected fully ranked HC3 binary graph results",
    )
    require(
        binary_permutation.groupby("endpoint")["significant_fdr"].sum().astype(int).to_dict()
        == expected_binary_counts,
        "Rank Freedman-Lane permutation check does not retain the binary graph conclusions",
    )
    require(len(continuous) == 11 * 15 and set(continuous["fdr_family"]) == {"all_15_global_metrics_within_outcome"}, "Continuous graph FDR family is inconsistent")
    figure6_source = read_csv(
        "results/analysis/graph_topology/figure_06_remaining_connectivity_graph_partial_correlations_source.csv"
    )
    figure6_inventory = read_csv("results/analysis/graph_topology/figure_06_panel_selection.csv")
    require(
        len(figure6_source) == 11 * 15
        and figure6_source["metric"].nunique() == 15
        and figure6_source["outcome"].nunique() == 11
        and not figure6_source.duplicated(["metric", "outcome"]).any(),
        "Figure 6 is not the complete fixed 15-descriptor by 11-outcome matrix",
    )
    require(
        figure6_source["display_selection"].str.contains("no result-based", case=False).all()
        and len(figure6_inventory) == 15
        and figure6_inventory["selection_rule"].str.contains("fixed independently", case=False).all(),
        "Figure 6 provenance does not exclude result-based display selection",
    )
    figure6_comparison = continuous.merge(
        figure6_source[["metric", "outcome", "rho_partial", "q_partial", "significant_partial_fdr"]],
        on=["metric", "outcome"],
        how="outer",
        validate="one_to_one",
        suffixes=("_analysis", "_figure"),
        indicator=True,
    )
    require(
        figure6_comparison["_merge"].eq("both").all()
        and np.allclose(
            figure6_comparison["rho_partial_analysis"],
            figure6_comparison["rho_partial_figure"],
            equal_nan=True,
        )
        and np.allclose(
            figure6_comparison["q_partial_analysis"],
            figure6_comparison["q_partial_figure"],
            equal_nan=True,
        )
        and figure6_comparison["significant_partial_fdr_analysis"].equals(
            figure6_comparison["significant_partial_fdr_figure"]
        ),
        "Figure 6 cells do not match the canonical continuous graph inference table",
    )
    require(len(nodal) == 11 * 7 * 191, "Nodal graph table must contain 11 outcomes x 7 metrics x 191 parcels")
    require(set(nodal["fdr_family"]) == {"all_7_metrics_x_191_parcels_within_outcome"}, "Nodal graph FDR family is not jointly defined")
    calculated = nodal.groupby("outcome")["significant_all_metrics"].sum().astype(int)
    reported = nodal_counts.set_index("outcome")["n_sig_parcel_metric"].astype(int)
    require(calculated.sort_index().equals(reported.sort_index()), "Nodal graph summary uses a different FDR family")


def verify_multivariate_and_prediction() -> None:
    joint = read_csv("results/analysis/multivariate/conditional_component_tests.csv")
    endpoint = read_csv("results/analysis/multivariate/endpoint_specific_conditional_tests.csv")
    require(len(joint) == 4 and len(endpoint) == 9, "Unexpected multivariate test families")
    assert_no_terms(joint, ("hub", "hvs", "graph", "loss", "delta"), "Multivariate joint tests")
    assert_no_terms(endpoint, ("hub", "hvs", "graph", "loss", "delta"), "Multivariate endpoint tests")
    expected_multivariate_method = {
        "Freedman-Lane outcome-residual permutation under reduced covariate model"
    }
    for frame, label in [(joint, "Multivariate joint tests"), (endpoint, "Multivariate endpoint tests")]:
        require(
            set(frame["permutation_method"]) == expected_multivariate_method
            and set(frame["n_permutations"]) == {5000}
            and set(frame["permutation_seed"]) == {42},
            f"{label} omits formal permutation provenance",
        )
    require(
        joint.loc[joint["domain"].eq("AAT_subtests"), "covariates"].str.contains("left_mean_chaco").all()
        and joint.loc[joint["domain"].eq("DemTect_subtests"), "covariates"].str.contains("global_mean_chaco").all(),
        "Joint multivariate provenance omits the mean ChaCo adjustment",
    )
    demtect_anatomical = joint.loc[
        joint["model"].eq("fs191_anatomical_7df_given_global_chaco")
    ]
    require(
        len(demtect_anatomical) == 1
        and int(demtect_anatomical.iloc[0]["n_predictors"]) == 7,
        "DemTect anatomical block must omit one reference group when conditioning on global mean ChaCo",
    )
    require(
        endpoint.loc[endpoint["domain"].eq("AAT_subtests"), "covariates"].str.contains("left_mean_chaco").all()
        and endpoint.loc[endpoint["domain"].eq("DemTect_subtests"), "covariates"].str.contains("global_mean_chaco").all(),
        "Endpoint multivariate provenance omits the mean ChaCo adjustment",
    )

    screen = read_csv("results/analysis/prediction/binary_model_screen_results.csv")
    reported = read_csv("results/analysis/prediction/reported_binary_model_details.csv")
    profile = read_csv("results/analysis/prediction/demtect_profile_model_summary.csv")
    profile_predictions = read_csv(
        "results/analysis/prediction/demtect_profile_cv_predictions_long.csv"
    )
    subtests = read_csv("results/analysis/prediction/reported_demtect_subtest_model_details.csv")
    table = read_csv("results/analysis/prediction/table_03_prediction_models.csv")
    for frame, label in [(screen, "Prediction screen"), (reported, "Reported prediction"), (profile, "DemTect profile prediction"), (subtests, "DemTect subtest prediction"), (table, "Table 3")]:
        assert_no_terms(frame, ("hub", "hvs", "graph", "loss", "winner", "selected best"), label)
    require(len(screen) == 18, "Prediction screen must compare 6 endpoints x 3 common feature sets")
    require(len(reported) == 6 and set(reported["feature_set"]) == {"clinical_lobe_roi_pca"}, "Reported binary endpoints do not use one common regional extension")
    require(len(profile) == 3 and len(subtests) == 5 and len(table) == 17, "Prediction result families are incomplete")
    require(list(table.columns) == ["Target", "Metric", "Clinical only", "Clinical + location", "Regional extension", "Difference vs clinical + location"], "Table 3 must use six common columns without panels")
    require(table["Metric"].isin(["AUC", "Balanced accuracy"]).sum() == 12
            and table["Target"].str.startswith("DemTect profile;").sum() == 2
            and table["Target"].str.startswith("DemTect individual subtests (range);").sum() == 3,
            "Table 3 outcome families are incomplete")
    for target, saved in screen.groupby("target_label"):
        source = saved.set_index("feature_set")
        for metric, field in [("AUC", "auc"), ("Balanced accuracy", "balanced_accuracy")]:
            selected = table[table["Target"].str.startswith(f"{target}: ") & table["Metric"].eq(metric)]
            require(len(selected) == 1, f"Missing Table 3 metric: {target}/{metric}")
            entry = selected.iloc[0]
            for column, feature in [("Clinical only", "clinical"), ("Clinical + location", "clinical_lobe"), ("Regional extension", "clinical_lobe_roi_pca")]:
                if column == "Clinical only" and metric != "AUC":
                    require(entry[column] == "—", "Unexpected clinical-only secondary metric")
                    continue
                r = source.loc[feature]
                expected = f"{r[field]:.3f}\n({r[field + '_ci_low']:.3f} to {r[field + '_ci_high']:.3f})"
                require(entry[column] == expected, f"Table 3 differs from saved estimates: {target}/{metric}/{column}")
            r = source.loc["clinical_lobe_roi_pca"]
            field = f"delta_{field}_vs_clinical_lobe"
            expected = f"{r[field]:.3f}\n({r[field + '_ci_low']:.3f} to {r[field + '_ci_high']:.3f})"
            require(entry["Difference vs clinical + location"] == expected, f"Table 3 paired interval mismatch: {target}/{metric}")
    require(table.equals(read_csv("outputs/tables/table_03_prediction_models.csv")), "Exported Table 3 differs from the analysis table")
    require(reported["exploratory"].astype(bool).all() and subtests["exploratory"].astype(bool).all(), "Prediction outputs must be labelled exploratory")
    require({"auc_ci_low", "auc_ci_high", "delta_auc_vs_clinical_lobe_ci_low", "delta_auc_vs_clinical_lobe_ci_high"}.issubset(reported.columns), "Binary prediction intervals are missing")
    required_profile_columns = {
        "summed_raw_component_score_c_index",
        "fold_mean_profile_benchmark_median_profile_r",
        "delta_median_profile_r_vs_fold_mean_profile_benchmark",
        "delta_median_profile_r_vs_fold_mean_profile_benchmark_ci_low",
        "delta_median_profile_r_vs_fold_mean_profile_benchmark_ci_high",
    }
    require(
        required_profile_columns.issubset(profile.columns),
        "DemTect profile outputs omit the aggregate-score clarification or mean-profile benchmark",
    )
    require(
        not any("total_score" in column for column in profile.columns),
        "DemTect raw subtest-score sum is mislabeled as the clinical total score",
    )
    benchmark_predictions = profile_predictions[
        profile_predictions["feature_set"].eq("fold_mean_profile_benchmark")
    ]
    require(
        len(benchmark_predictions) == 132 * 5
        and benchmark_predictions["subject_id"].nunique() == 132,
        "Fold-contained DemTect mean-profile benchmark predictions are incomplete",
    )
    profile_rows = table[table["Target"].str.startswith("DemTect profile;")].set_index("Metric")
    require(set(profile_rows.index) == {"Median within-patient profile r", "Summed raw subtest-score C-index"},
            "Table 3 must distinguish the profile from the raw subtest-score sum")
    regional_profile = profile[profile["feature_set"].eq("clinical_lobe_roi_pca")].iloc[0]

    def check_interval(display: str, row: pd.Series, field: str) -> None:
        expected = f"{row[field]:.3f}\n({row[field + '_ci_low']:.3f} to {row[field + '_ci_high']:.3f})"
        require(display == expected, f"Table 3 continuous interval mismatch: {field}")

    for metric, field, delta in [
        ("Median within-patient profile r", "median_profile_r", "delta_median_profile_r_vs_clinical_lobe"),
        ("Summed raw subtest-score C-index", "summed_raw_component_score_c_index", "delta_summed_raw_component_score_c_index_vs_clinical_lobe"),
    ]:
        row = profile_rows.loc[metric]
        for column, feature in [("Clinical + location", "clinical_lobe"), ("Regional extension", "clinical_lobe_roi_pca")]:
            check_interval(row[column], profile[profile["feature_set"].eq(feature)].iloc[0], field)
        check_interval(row["Difference vs clinical + location"], regional_profile, delta)
    # Benchmark values moved to the generated note, preserving both model comparisons.
    from docx import Document
    document = Document(ROOT / "outputs/tables/table_03_prediction_models.docx")
    require(len(document.tables) == 1, "Table 3 Word export must contain one table")
    require(len(document.tables[0].rows) == 18 and len(document.tables[0].columns) == 6,
            "Table 3 Word export omits rows or columns")
    note = " ".join(paragraph.text for paragraph in document.paragraphs)
    baseline_profile = profile[profile["feature_set"].eq("clinical_lobe")].iloc[0]
    for saved, field in [
        (regional_profile, "fold_mean_profile_benchmark_median_profile_r"),
        (regional_profile, "delta_median_profile_r_vs_fold_mean_profile_benchmark"),
        (baseline_profile, "delta_median_profile_r_vs_fold_mean_profile_benchmark"),
    ]:
        expected = f"{saved[field]:.3f} ({saved[field + '_ci_low']:.3f} to {saved[field + '_ci_high']:.3f})"
        require(expected in note, "Table 3 note omits or changes a training-mean benchmark estimate")
    require("multi-output RidgeCV" in note, "Table 3 note omits the continuous-outcome model")
    ranges = table[table["Target"].str.startswith("DemTect individual subtests (range);")].set_index("Metric")
    for metric, field in [("C-index", "c_index"), ("Spearman rho", "spearman"), ("R²", "r2")]:
        require(ranges.loc[metric, "Regional extension"] == f"{subtests[field].min():.3f} to {subtests[field].max():.3f}",
                f"Table 3 subtest range mismatch: {metric}")
    delta = subtests["delta_c_index_vs_clinical_lobe"]
    require(ranges.loc["C-index", "Difference vs clinical + location"] == f"{delta.min():.3f} to {delta.max():.3f}",
            "Table 3 subtest difference range mismatch")


def verify_qol() -> None:
    correlations = read_csv("results/analysis/qol/qol_neuropsych_correlations.csv")
    groups = read_csv("results/analysis/qol/qol_impairment_group_differences.csv")
    structural = read_csv("results/analysis/qol/qol_structural_disconnection_correlations.csv")
    heatmap = read_csv("results/analysis/qol/qol_heatmap_source.csv")
    require(len(correlations) == 12 * 19, "QoL continuous screen must contain 12 neuropsychological measures x 19 QoL endpoints")
    require(correlations["neuropsych_measure"].nunique() == 12 and correlations["qol_measure"].nunique() == 19, "QoL continuous screen is incomplete")
    require("AAT_total" in set(correlations["neuropsych_measure"]), "QoL screen omits AAT total")
    require(len(groups) == 6 * 19 and groups["impairment_measure"].nunique() == 6, "QoL Mann-Whitney/Cliff screen is incomplete")
    require({"mannwhitney_u", "cliffs_delta_raw", "cliffs_delta_worse_status_aligned"}.issubset(groups.columns), "QoL group-effect statistics are missing")
    require(len(structural) == 57 and len(heatmap) == 75, "QoL structural/display families are incomplete")
    assert_no_terms(structural, ("hub", "hvs", "connectivity_loss", "efficiency_loss", "_delta"), "QoL structural tests")
    graph = structural[structural["feature_family"].eq("nemosc_graph_descriptor")]
    require(len(graph) == 45 and graph["structural_feature"].nunique() == 15, "QoL graph family must contain 15 descriptors x 3 outcomes")
    require(graph["q_fdr_45_graph_qol_tests"].notna().all(), "QoL graph family FDR is incomplete")
    fixed = structural[structural["feature_family"].eq("fixed_regional_chaco")]
    require(len(fixed) == 6 and fixed["q_fdr_6_fixed_regional_qol_tests"].notna().all(), "QoL fixed regional family is incomplete")
    joint = read_csv("results/analysis/qol/qol_structural_joint_51_sensitivity.csv")
    joint_source = structural[structural["feature_family"].isin(
        ["fixed_regional_chaco", "nemosc_graph_descriptor"]
    )].reset_index(drop=True)
    require(len(joint) == 51 and joint[joint_source.columns].equals(joint_source),
            "Joint QoL sensitivity must retain all six regional and 45 graph tests")
    # Independent SciPy implementation checks the BH correction in the writer.
    from scipy.stats import false_discovery_control
    require(np.allclose(joint["q_fdr_joint_51"], false_discovery_control(joint["p"], method="bh")),
            "Joint structural QoL correction does not match BH across all 51 tests")
    require(joint["significant_joint_51"].eq(joint["q_fdr_joint_51"] < 0.05).all(),
            "Joint structural QoL significance flags are inconsistent")
    selected = structural[structural["structural_feature"].str.contains("TFNBS", case=False, na=False)]
    require(len(selected) == 6 and selected["same_sample_outcome_selected_feature"].astype(bool).all(), "TFNBS QoL edge summaries are not marked non-independent")
    require(selected["selection_note"].str.contains("non-independent", case=False, na=False).all(), "TFNBS QoL selection note is incomplete")
    require(selected["feature_family"].eq("outcome_selected_tfnbs").all(), "TFNBS QoL rows are not separated from fixed regional inference")
    require((~selected["inference_eligible"].astype(bool)).all(), "TFNBS QoL rows are incorrectly inference-eligible")
    require(selected[["p", "q_fdr_6_fixed_regional_qol_tests", "q_fdr_45_graph_qol_tests", "q_fdr_family"]].isna().all().all(), "TFNBS QoL rows must not contain p or q values")
    require((~selected["fdr_significant"].astype(bool)).all(), "TFNBS QoL rows must not receive significance markers")
    selected_heatmap = heatmap[heatmap["row_label"].str.contains("TFNBS", case=False, na=False)]
    require(len(selected_heatmap) == 6 and (~selected_heatmap["inference_eligible"].astype(bool)).all(), "Figure 7 TFNBS cells are incorrectly inference-eligible")
    require((~selected_heatmap["fdr_significant"].astype(bool)).all(), "Figure 7 must not star outcome-selected TFNBS rows")


def verify_manuscript_outputs() -> None:
    expected_figures = {f"figure_{number:02d}_{name}.png" for number, name in [
        (1, "cohort_outcomes"),
        (2, "regional_disconnection_patterns"),
        (3, "global_subtest_disconnection_associations"),
        (4, "subtest_tfnbs_edge_configurations"),
        (5, "graph_topological_network_impact"),
        (6, "graph_metric_subtest_associations"),
        (7, "qol_associations"),
    ]}
    observed_figures = {path.name for path in (ROOT / "outputs/figures").glob("*.png")}
    require(observed_figures == expected_figures, f"Figure set mismatch: expected={sorted(expected_figures)}, observed={sorted(observed_figures)}")
    for path in sorted((ROOT / "outputs/figures").glob("*.png")):
        with Image.open(path) as image:
            require(image.width >= 800 and image.height >= 500, f"Figure is unexpectedly small: {path.name} ({image.size})")
            extrema = np.asarray(image.convert("RGB")).ptp(axis=(0, 1))
            require(bool((extrema > 10).any()), f"Figure appears blank: {path.name}")

    expected_tables = {"table_01_cohort_characteristics.csv", "table_02_tfnbs_subtest_summary.csv", "table_03_prediction_models.csv"}
    observed_tables = {path.name for path in (ROOT / "outputs/tables").glob("*.csv")}
    require(observed_tables == expected_tables, f"Table set mismatch: expected={sorted(expected_tables)}, observed={sorted(observed_tables)}")
    table3 = pd.read_csv(ROOT / "outputs/tables/table_03_prediction_models.csv")
    assert_no_terms(table3, ("hub", "hvs", "graph", "loss", "winner"), "Manuscript Table 3")

    provenance = read_csv("metadata/figure_generation_provenance.csv")
    require(set(provenance["figure"]) == expected_figures, "Figure provenance is incomplete")
    require(~provenance["method"].str.contains("copied|static", case=False, regex=True).any(), "Figure provenance contains copied/static assets")

    checksum_path = ROOT / "metadata/output_checksums_sha256.csv"
    require(checksum_path.exists(), "Missing release checksum manifest")
    with checksum_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    expected_paths = {f"outputs/figures/{name}" for name in expected_figures} | {f"outputs/tables/{name}" for name in expected_tables}
    listed_paths = {row["relative_path"] for row in rows}
    require(listed_paths == expected_paths, "Release checksum manifest does not exactly cover the 7 figures and 3 tables")
    for row in rows:
        path = ROOT / row["relative_path"]
        require(path.exists(), f"Checksum manifest lists a missing output: {row['relative_path']}")
        require(sha256_file(path) == row["sha256"], f"Checksum mismatch: {row['relative_path']}")


def main() -> None:
    cohort = read_csv("results/analysis/cohort/merged_cohort.csv")
    require(len(cohort) == 163 and cohort["subject_id"].is_unique, "Cohort must contain 163 unique subjects")
    verify_nemo_encoding()
    verify_regional_statistics()
    verify_tfnbs()
    verify_graph_statistics()
    verify_multivariate_and_prediction()
    verify_qol()
    verify_manuscript_outputs()
    print("OK: verified corrected scientific semantics, complete results, generated displays, and release checksums")


if __name__ == "__main__":
    main()
