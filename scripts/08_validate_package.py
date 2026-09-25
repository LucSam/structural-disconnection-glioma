#!/usr/bin/env python3
"""Check result provenance, correction families and document exports."""
import importlib.util
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from scipy.stats import spearmanr

from _common import ROOT, bh_fdr, sha256, holm


def main():
    checks = {}

    def check(name, condition):
        checks[name] = bool(condition)
        if not condition:
            raise AssertionError(name)

    manifest = json.loads((ROOT / 'metadata/input_manifest.json').read_text())
    for row in manifest:
        check(f"snapshot unchanged: {row['destination']}",
              sha256(ROOT / row['destination']) == row['destination_sha256'])
    original = ROOT.parent / 'manuscript-data'
    if original.exists():
        for row in manifest:
            check(f"original unchanged: {row['source']}",
                  sha256(original / row['source']) == row['source_sha256'])

    reproduced = pd.read_csv(ROOT / 'results/original_joint_reproduction.csv')
    check('all four original joint models reproduced', len(reproduced) == 4 and reproduced.passed.all())
    models = pd.read_csv(ROOT / 'results/staged_models.csv')
    check('six model comparisons, 19,999 permutations each',
          len(models) == 6 and models.n_permutations.eq(19999).all())
    check('FDR across the six model comparisons',
          np.allclose(models.q_fdr_six_models, bh_fdr(models.p_permutation), atol=1e-14))
    profile = pd.read_csv(ROOT / 'results/profile/profile_omnibus_tests.csv')
    pairs = pd.read_csv(ROOT / 'results/profile/pairwise_profile_tests.csv')
    check('five profile tests corrected together',
          len(profile) == 5 and np.allclose(profile.p_holm_5, holm(profile.p_permutation)))
    check('16 pairwise contrasts corrected together',
          len(pairs) == 16 and np.allclose(pairs.p_holm_16, holm(pairs.p_permutation)))
    qol = pd.read_csv(ROOT / 'results/qol_correlations.csv')
    reference = pd.read_csv(ROOT / 'data/reference/qol_neuropsych_correlations.csv')
    check('222 estimable QoL tests retained from 228 pairs', len(qol) == 228 and qol.p_value.notna().sum() == 222)
    keys = ['neuropsych_measure', 'qol_measure']
    q = qol.set_index(keys).sort_index()
    old = reference.set_index(keys).sort_index()
    check('all QoL correlations reproduce the original family',
          q.index.equals(old.index) and all(np.allclose(q[c], old[c], equal_nan=True, atol=1e-12)
                                          for c in ['n', 'spearman_rho', 'p_value', 'q_fdr_all_neuropsych_qol_tests']))
    check('QoL FDR covers the full family',
          np.allclose(qol.q_fdr_all_neuropsych_qol_tests, bh_fdr(qol.p_value), equal_nan=True))
    thirds = pd.read_csv(ROOT / 'results/burden_tertiles.csv')
    groups = json.loads((ROOT / 'metadata/clinical_definitions.json').read_text())
    for battery, g in groups.items():
        sub = thirds[thirds.battery.eq(battery)]
        high = sub[sub.tertile.eq(3)].iloc[0]
        check(f'{battery} description and model samples agree',
              sub.n.sum() == g['n'] and sub.preserved.sum() == g['n_preserved'] and
              high.n == g['n_high_burden'] and high.preserved == g['n_high_burden_preserved'] and
              models[models.battery.eq(battery)].n.eq(g['n']).all())

    tf = pd.read_csv(ROOT / 'results/tfnbs_count_verification.csv')
    check('all 36 TFNBS model/correction counts verified', len(tf) == 36)
    source_counts = pd.read_csv(ROOT / 'data/reference/tfnbs_summary.csv')
    table3 = pd.read_csv(ROOT / 'results/connection_counts.csv')
    # The dedicated anatomy step performs the edge-level recount; here verify
    # that the report build has not changed the resulting count matrix.
    check('nine TFNBS subtests retained', len(table3) == len(source_counts) == 9)
    for r in table3.itertuples(index=False, name=None):
        battery, subtest, *counts = r
        expected = [int(tf[tf.battery.eq(battery) & tf.subtest.eq(subtest) &
                           tf.model.eq(model) & tf.correction.eq(correction)].n_edges.iloc[0])
                    for model in ['clinical', 'lobe'] for correction in ['edge FWE', 'across subtests']]
        check(f'TF NBS table values: {battery}, {subtest}', counts == expected)

    path = ROOT / 'scripts/07_build_manuscript.py'
    spec = importlib.util.spec_from_file_location('build_report', path)
    builder = importlib.util.module_from_spec(spec); spec.loader.exec_module(builder)
    context, model_frame, definitions = builder.context_from_results()
    table_specs = [builder.statistical_overview(), *builder.manuscript_tables(model_frame, definitions), *builder.supplementary_tables()]
    doc = Document(ROOT / 'manuscript/icons_gp_manuscript_v2.docx')
    check('four main and three supplementary tables; five main and five supplementary figures', len(doc.tables) == 7 and len(doc.inline_shapes) == 10)
    for spec, word_table in zip(table_specs, doc.tables, strict=True):
        df = spec['frame']
        check(f"table dimensions: {spec['stem']}", len(word_table.rows) == len(df) + 1 and len(word_table.columns) == len(df.columns))
        for values, word_row in zip(df.itertuples(index=False, name=None), word_table.rows[1:], strict=True):
            for j, (value, cell) in enumerate(zip(values, word_row.cells, strict=True)):
                expected = str(value)
                check(f"Word cell: {spec['stem']} / {values[0]} / {j} / {values[1]}", cell.text == expected)
    build = json.loads((ROOT / 'metadata/manuscript_build.json').read_text())
    check('manuscript values reflect current results', build['context'] == context)
    check('manuscript template hash is current', build['template_sha256'] == sha256(ROOT / 'manuscript/manuscript_template.md'))
    check('all result hashes match the manuscript build', all(sha256(ROOT / p) == h for p, h in build['result_hashes'].items()))
    check('embedded figures match the manuscript build', all(sha256(ROOT / p) == h for p,h in build['figure_hashes'].items()))
    text = '\n'.join(p.text for p in doc.paragraphs)
    check('no unresolved manuscript tokens', '{{' not in text and '}}' not in text)
    check('no study IDs or local user paths in manuscript', not re.search(r'\bP\d{3}\b|/Users/|subject_id', text))
    for folder in ['results', 'outputs/tables']:
        for file in (ROOT / folder).rglob('*.csv'):
            check(f'no individual identifiers: {file.relative_to(ROOT)}',
                  not re.search(r'\bP\d{3}\b|\bsubject_id\b|/Users/', file.read_text()))
    overlap = json.loads((ROOT / 'metadata/lesion_overlap_provenance.json').read_text())
    check('cohort overlap uses 163 masks and unchanged aggregate image',
          overlap['n'] == 163 and sha256(ROOT / overlap['destination']) == overlap['destination_sha256'])
    display = pd.read_csv(ROOT / 'data/private/display_participants.csv')
    aat = display[display.battery.eq('AAT')]
    check('lowest AAT margin matches the four-subtest definition',
          np.array_equal(aat.minimum_aat_margin.ge(0), aat.preserved))
    effects = pd.read_csv(ROOT / 'results/subtest_pattern_partial_r2.csv')
    check('nine descriptive subtest contributions from joint model samples',
          len(effects) == 9 and effects.partial_r2.between(0,1).all() and
          effects[effects.battery.eq('AAT')].n.eq(134).all() and
          effects[effects.battery.eq('DemTect')].n.eq(132).all())
    audit = json.loads((ROOT / 'metadata/hemisphere_audit.json').read_text())
    check('hemisphere audit private input hashes unchanged',
          all(sha256(ROOT / p) == h for p, h in audit['private_input_hashes'].items()))
    check('audit reproduces original adjusted TFNBS t statistics',
          all(m['max_abs_original_t_reproduction_error'] < 2e-5 for m in audit['summaries']))
    hemi = pd.read_csv(ROOT / 'results/hemisphere_audit/edge_hemisphere_counts.csv')
    naming = hemi[hemi.subtest.eq('Naming') & hemi.correction.eq('across subtests')]
    check('Naming endpoint hemisphere counts',
          naming.set_index('hemisphere_pair').n_edges.to_dict() == {'L–L': 4, 'L–M': 1, 'L–R': 141, 'M–R': 1, 'R–R': 5})
    token = hemi[hemi.subtest.eq('Token Test') & hemi.correction.eq('across subtests')]
    check('all 21 Token Test retained edges have left endpoints',
          token.set_index('hemisphere_pair').n_edges.to_dict() == {'L–L': 21})
    support = pd.read_csv(ROOT / 'results/hemisphere_audit/edge_support_and_influence.csv')
    check('all retained edges have left-lesion support and positive diagnostic associations',
          len(support) == 173 and support.n_nonzero_left_lesions.gt(0).all() and
          support.partial_deficit_rho_left_lesions_only.gt(0).all())
    check('AAT burden includes 82 cerebral and 10 cerebellar left parcels',
          audit['burden']['left_cerebral_parcels'] == 82 and
          audit['burden']['left_cerebellar_parcels'] == 10 and
          audit['burden']['high_tertile_membership_changes'] == 0)
    new = pd.read_csv(ROOT / 'results/tfnbs_hemisphere/edge_statistics.csv.gz')
    new_counts = pd.read_csv(ROOT / 'results/tfnbs_hemisphere/counts.csv')
    comparison = pd.read_csv(ROOT / 'results/connection_model_comparison.csv')
    check('all nine hemisphere-adjusted endpoints and all possible edges exported',
          len(new_counts) == len(comparison) == 9 and len(new) == 9*191*190//2)
    check('new adjusted t statistics independently reproduced', new_counts.max_abs_t_reproduction_error.lt(3e-5).all())
    for row in new_counts.itertuples():
        d = new[new.subtest_label.eq(row.subtest)]
        multiplier = 4 if row.battery == 'AAT' else 5
        check(f'new TFNBS correction and counts: {row.subtest}',
              np.allclose(d.p_battery_bonferroni, np.minimum(1, d.p_fwe_edges*multiplier)) and
              d.significant_fwe_edges.sum() == row.edge_FWE and
              d.significant_battery_bonferroni.sum() == row.across_subtests)
        comparison_row = comparison[comparison.subtest.eq(row.subtest)].iloc[0]
        old_row = table3[table3.Subtest.eq(row.subtest)].iloc[0]
        check(f'comparison table matches both source models: {row.subtest}',
              comparison_row.n == row.n and
              comparison_row.hemisphere_edge_FWE == row.edge_FWE and
              comparison_row.hemisphere_across_subtests == row.across_subtests and
              comparison_row.lobe_edge_FWE == old_row['Location-adjusted: edge FWE'] and
              comparison_row.lobe_across_subtests == old_row['Location-adjusted: across subtests'])
    tfnbs = json.loads((ROOT / 'metadata/hemisphere_tfnbs.json').read_text())
    check('new TFNBS uses 5000 shuffles and unchanged enhancement parameters',
          tfnbs['parameters']['nshuffles'] == 5000 and tfnbs['parameters']['tfce_dh'] == .1 and
          tfnbs['parameters']['tfce_e'] == .4 and tfnbs['parameters']['tfce_h'] == 3)
    for run in (ROOT / 'data/private/tfnbs_hemisphere/runs').glob('*/run_manifest.json'):
        m = json.loads(run.read_text())
        check(f'new TFNBS raw output hashes: {run.parent.name}',
              all(sha256(run.parent / f) == h for f,h in m['output_hashes'].items()))
        check(f'new TFNBS inputs match completed run: {run.parent.name}',
              m['signature']['edge_input_sha256'] == sha256(ROOT / 'data/private/tfnbs_hemisphere/chacoconn_edges.npz') and
              m['signature']['subjects_sha256'] == sha256(ROOT / 'data/private/tfnbs_hemisphere' / f'{run.parent.name}_subjects.csv') and
              m['signature']['parameters'] == tfnbs['parameters'])
    high = pd.read_csv(ROOT / 'results/high_burden_group_summary.csv')
    maps = pd.read_csv(ROOT / 'results/high_burden_parcel_means.csv')
    check('four descriptive maps each include all 191 parcels',
          len(high) == 4 and maps.groupby(['battery','group']).size().eq(191).all())
    for battery in ['AAT','DemTect']:
        d = high[high.battery.eq(battery)]
        check(f'high-burden map sample matches existing definition: {battery}',
              d.n.sum() == groups[battery]['n_high_burden'] and
              d.loc[d.preserved,'n'].sum() == groups[battery]['n_high_burden_preserved'])
    smoothing = pd.read_csv(ROOT / 'results/qol_display_smoothers.csv')
    check('QoL smoothers are finite and limited to the two displayed relationships',
          set(smoothing.battery) == {'AAT','DemTect'} and np.isfinite(smoothing[['score','fitted_communication']]).all().all())
    region = pd.read_csv(ROOT / 'results/regional_display_correlations.csv')
    check('eleven complete, unthresholded regional coefficient maps',
          len(region) == 11*191 and region.groupby('outcome').size().eq(191).all() and
          region.deficit_rho.between(-1,1).all() and 'p' not in region and 'q' not in region)
    regional_metadata = json.loads((ROOT / 'metadata/regional_displays.json').read_text())
    check('regional DemTect map reproduces saved hemisphere-adjusted coefficients',
          regional_metadata['max_error_saved_DemTect_map'] < 1e-10)
    differences = pd.read_csv(ROOT / 'results/high_burden_regional_differences.csv')
    check('group difference maps contain 191 parcels per battery and correct sign',
          len(differences) == 382 and np.allclose(differences.difference_percentage_points,
          100*(differences['Deficit']-differences['No deficit'])))
    displayed_edges = json.loads((ROOT / 'metadata/connection_figure_counts.json').read_text())
    check('main and sensitivity edge displays are uncapped and use the same criterion',
          len(displayed_edges) == 8 and all(r['n_drawn'] == r['n_qualifying'] and r['cap'] is None and
          r['criterion'] == 'significant_battery_bonferroni' for r in displayed_edges))
    for r in displayed_edges:
        column = 'lobe_across_subtests' if r['figure'].startswith('figure_S5') else 'hemisphere_across_subtests'
        expected = comparison.set_index('subtest').loc[r['subtest'],column]
        check(f"all qualifying displayed edges: {r['figure']} / {r['subtest']}", r['n_drawn'] == expected)
    check('supplement source current and embedded after main references',
          build['supplement_template_sha256'] == sha256(ROOT / 'manuscript/supplement_template.md') and
          text.index('Supplementary material') > text.index('References') and
          'Supplementary references' in text)
    check('full profile and pairwise results retained in supplement',
          'Table S1. Joint profile checks.' in text and 'Table S2. Pairwise subtest profile contrasts.' in text)
    network_meta = json.loads((ROOT / 'metadata/network_context.json').read_text())
    check('structural context inputs match current files',
          all(sha256(ROOT / p) == h for p,h in network_meta['input_hashes'].items()))
    source = np.load(ROOT / 'data/private/tfnbs_hemisphere/chacoconn_edges.npz')
    edge_values = source['values'].astype(float)
    coverage = pd.read_csv(ROOT / 'results/cohort_edge_coverage.csv')
    grouped = pd.read_csv(ROOT / 'results/cohort_edge_group_coverage.csv')
    check('coverage includes every unselected atlas pair and zeros',
          len(coverage) == 18145 and len(grouped) == 36 and
          grouped.n_edge_pairs.sum() == 18145 and
          np.allclose(coverage.mean_chacoconn,edge_values.mean(axis=0)) and
          np.array_equal(coverage.n_nonzero,(edge_values>0).sum(axis=0)) and
          np.isclose(np.average(grouped.mean_chacoconn,weights=grouped.n_edge_pairs),edge_values.mean()))
    edge_means = pd.read_csv(ROOT / 'results/high_burden_edge_means.csv.gz')
    group_means = pd.read_csv(ROOT / 'results/high_burden_edge_group_means.csv')
    grouping = pd.read_csv(ROOT / 'results/structural_parcel_groups.csv').set_index('roi_index')
    check('structural grouping contains only anatomical assignments and side labels',
          len(grouping) == 191 and grouping.anatomical_group.nunique() == 8 and
          set(grouping.columns) == {'roi_name','anatomical_group','side'})
    for (battery,group),d in group_means.groupby(['battery','group']):
        m = edge_means[edge_means.battery.eq(battery)&edge_means.group.eq(group)]
        cases = display[display.battery.eq(battery)&display.tertile.eq(3)&display.preserved.eq(group=='No deficit')]
        indices = np.flatnonzero(np.isin(source['subject_ids'].astype(str),cases.subject_id))
        check(f'all high-burden edges reproduce source patients: {battery}, {group}',
              len(m) == 18145 and len(d) == 36 and len(indices) == d.n.iloc[0] and
              np.allclose(m.mean_chacoconn,edge_values[indices].mean(axis=0)))
        for r in d.itertuples():
            gi = grouping.loc[m.roi_i,'anatomical_group'].to_numpy()
            gj = grouping.loc[m.roi_j,'anatomical_group'].to_numpy()
            mask = ((gi==r.group_i)&(gj==r.group_j)) | ((gi==r.group_j)&(gj==r.group_i))
            check(f'anatomical edge mean: {battery}, {group}, {r.group_i}, {r.group_j}',
                  mask.sum() == r.n_edge_pairs and np.isclose(m.loc[mask,'mean_chacoconn'].mean(),r.mean_chacoconn))
    edge_diff = pd.read_csv(ROOT / 'results/high_burden_edge_group_differences.csv')
    check('edge differences use deficit minus no deficit without tests',
          len(edge_diff) == 72 and np.allclose(edge_diff.difference_percentage_points,
          100*(edge_diff['Deficit']-edge_diff['No deficit'])) and 'p' not in edge_diff and 'q' not in edge_diff)
    network = pd.read_csv(ROOT / 'results/network_qol_correlations.csv')
    private_qol = pd.read_csv(ROOT / 'data/private/network_qol_participants.csv').set_index('subject_id')
    selected = pd.read_csv(ROOT / 'results/qol_connection_sets.csv')
    for battery,n_edges in [('AAT',2930),('DemTect',404)]:
        expected_pairs = set(map(tuple,new.loc[new.battery.eq(battery)&new.significant_battery_bonferroni,['roi_i','roi_j']].to_numpy()))
        observed_pairs = set(map(tuple,selected.loc[selected.battery.eq(battery),['roi_i','roi_j']].to_numpy()))
        check(f'QoL edge set is the deduplicated main-model union: {battery}',
              expected_pairs == observed_pairs and len(observed_pairs) == n_edges)
        mask = np.array([(int(i),int(j)) in observed_pairs for i,j in zip(source['roi_i'],source['roi_j'])])
        check(f'selected-edge patient means reproduce raw snapshot: {battery}',
              np.allclose(private_qol.loc[source['subject_ids'].astype(str),f'{battery}_edges'],edge_values[:,mask].mean(axis=1)))
    check('all structural QoL displays are descriptive, with no new p or q family',
          len(network) == 12 and 'p' not in network and 'q' not in network)
    for r in network.itertuples():
        d = private_qol[[r.feature,r.qol_measure]].dropna()
        check(f'network–QoL coefficient: {r.feature}, {r.qol_label}',
              len(d) == r.n and np.isclose(spearmanr(d.iloc[:,0],d.iloc[:,1]).statistic,r.rho))
    prior = pd.read_csv(ROOT / 'results/prior_regional_qol_families.csv')
    joint = pd.read_csv(ROOT / 'data/reference/qol_structural_joint_51_sensitivity.csv')
    check('historical six- and 51-test corrections retained',
          len(prior) == 6 and len(joint) == 51 and
          np.allclose(prior.q_six,bh_fdr(prior.p)) and
          np.allclose(joint.q_fdr_joint_51,bh_fdr(joint.p)) and
          prior.q_six.lt(.05).all() and prior.q_joint_51.ge(.05).all())
    check('historical structural QoL table embedded in supplement',
          'Table S3. Regional disconnection–QoL correlations and correction families.' in text)
    references = json.loads((ROOT / 'manuscript/references_used.json').read_text())
    check('all references have a year and complete bibliographic ending',
          all(re.search(r'\(\d{4}\)', r['text']) and
              ('https://doi.org/' in r['text'] or r['key'] == 'Huber 1984') for r in references))
    for stem, _, _ in builder.FIGURES + builder.SUPPLEMENT_FIGURES:
        check(f'PNG figure available: {stem}',
              (ROOT / 'outputs/figures' / f'{stem}.png').is_file())
    check('figure outputs use PNG only',not list((ROOT/'outputs').rglob('*.svg')))
    headings=[p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    result_headings=[h for h in headings if re.match(r'3\.[1-5] ',h)]
    check('results follow cohort, groups, model, anatomy and patient reports',
          result_headings == ['3.1 Cohort and disconnection coverage',
          '3.2 Patients with and without deficits at high burden',
          '3.3 Associations with disconnection burden and pattern',
          '3.4 Where disconnection was associated with performance',
          '3.5 Patient-reported function'])
    check('abstract at most 250 words',build['abstract_words'] <= 250)
    check('no gating ambiguity in main overview',
          'gated by parent test' not in builder.statistical_overview()['frame'].to_string())
    schema=json.loads((ROOT/'metadata/nemo_schema.json').read_text())
    check('NeMo diagram is identified as synthetic and its example fractions agree',
          'Synthetic' in schema['type'] and
          np.allclose(schema['regional_percent'], [100/3,37.5,0,25,31.25,25]) and
          np.isclose(schema['pair_percent']['B–E'],100) and
          np.isclose(schema['pair_percent']['A–D'],75) and
          np.isclose(schema['pair_percent']['A–E'],25) and
          np.isclose(schema['pair_percent']['B–F'],50) and
          sum(schema['incident_affected_weight'])==2*schema['total_affected_weight']==20 and
          sum(schema['incident_reference_weight'])==2*schema['total_reference_weight']==72 and
          all(np.isclose(e['percent'],100*e['affected_weight']/e['reference_weight']) and
              (e['percent']==0 or np.isclose(e['line_width_points'],.08*e['percent']))
              for e in schema['connections']))
    overlap=json.loads((ROOT/'metadata/battery_overlap.json').read_text())
    paired=display.pivot(index='subject_id',columns='battery',values=['burden','preserved','tertile']).dropna()
    check('battery overlap uses paired complete cases',len(paired)==overlap['n_both_complete']==104)
    check('battery status concordance agrees with complete-case labels',
          paired['preserved']['AAT'].eq(paired['preserved']['DemTect']).sum()==overlap['n_same_status']==80)
    high=paired[paired['tertile']['AAT'].eq(3)&paired['tertile']['DemTect'].eq(3)]
    check('high-burden overlap uses original battery-specific thirds',
          len(high)==overlap['n_both_high']==29 and
          high['preserved']['AAT'].eq(high['preserved']['DemTect']).sum()==overlap['n_high_same_status']==24)
    check('reported burden correlation uses the paired sample',
          np.isclose(spearmanr(paired['burden']['AAT'],paired['burden']['DemTect']).statistic,
                     overlap['burden_spearman_rho_paired']) and context['overlap_burden_rho']=='0.885')
    overlap_cells=pd.read_csv(ROOT/'results/battery_status_overlap.csv')
    check('aggregate overlap table preserves all cells',
          overlap_cells.groupby('sample').n.sum().to_dict()=={'Both complete batteries':104,'Both highest thirds':29})
    alternative=json.loads((ROOT/'metadata/alternative_circle_counts.json').read_text())
    check('alternative circle uses current main-model results',
          alternative['source_sha256']==sha256(ROOT/'results/tfnbs_hemisphere/edge_statistics.csv.gz'))
    node_counts=pd.read_csv(ROOT/'outputs/alternatives/figure_04_alternative_node_counts.csv')
    check('alternative circle retains a fixed node order across panels',
          node_counts.groupby('roi_index').angle_radians.nunique().eq(1).all())
    for row in alternative['panels']:
        n=int(new.loc[new.subtest_label.eq(row['subtest'])&new.significant_battery_bonferroni].shape[0])
        local=node_counts[node_counts.subtest.eq(row['subtest'])]
        check(f"uncapped alternative circle and incident counts: {row['subtest']}",
              row['n_drawn']==row['n_qualifying']==n and row['cap'] is None and
              len(local)==191 and local.displayed_connection_count.sum()==2*n)
    result = {'passed': all(checks.values()), 'number_of_checks': len(checks), 'checks': checks,
              'note': 'These checks complement the permutation tests and visual review; they do not replace scientific review.'}
    (ROOT / 'metadata/validation.json').write_text(json.dumps(result, indent=2) + '\n')
    print(f'Passed {len(checks)} package checks.')


if __name__ == '__main__':
    main()
